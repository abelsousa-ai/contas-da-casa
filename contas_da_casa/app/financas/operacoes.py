"""Operações feitas a partir da interface."""
import datetime as dt, io
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font
from .normalizar import norm_desc


def _agora():
    return dt.datetime.now().isoformat(timespec='seconds')


def listas(con):
    cats = [dict(id=i, nome=n, tipo=t, activa=a) for i, n, t, a in
            con.execute('SELECT id, nome, tipo, activa FROM categorias ORDER BY nome')]
    subs = [dict(id=i, categoria_id=c, nome=n, na_lista=l, activa=a) for i, c, n, l, a in
            con.execute('SELECT id, categoria_id, nome, na_lista, activa FROM subcategorias ORDER BY nome')]
    pessoas = [dict(id=i, nome=n, activa=a) for i, n, a in con.execute('SELECT id, nome, activa FROM pessoas ORDER BY id')]
    contas = [dict(id=i, codigo=c, nome=n, so_historico=h) for i, c, n, h in
              con.execute('SELECT id, codigo, nome, so_historico FROM contas ORDER BY id')]
    return dict(categorias=cats, subcategorias=subs, pessoas=pessoas, contas=contas)


def sugestao(con, descricao_norm):
    """Classificação mais frequente desta descrição (só 2026 em diante), para pré-preencher."""
    r = con.execute("SELECT valor FROM meta WHERE chave = 'aberto_de'").fetchone()
    desde = r[0] if r else '1900-01-01'
    hist = con.execute('''SELECT categoria_id, subcategoria_id, pessoa_id FROM movimentos
                          WHERE descricao_norm = ? AND estado = 'ok' AND origem != 'previsto' AND data_lanc >= ?''',
                       (descricao_norm, desde)).fetchall()
    if not hist:
        return None
    (cat, sub, pes), n = Counter(hist).most_common(1)[0]
    return dict(categoria_id=cat, subcategoria_id=sub, pessoa_id=pes, vezes=n, total=len(hist))


def classificar(con, mov_id, categoria_id, subcategoria_id, pessoa_id, mesma_descricao=False):
    """Classificação pelo utilizador. Propaga às taxas desta compra e, se pedido, aos outros
    pendentes com a mesma descrição. Devolve o número de movimentos alterados."""
    m = con.execute('SELECT descricao_norm, origem FROM movimentos WHERE id = ?', (mov_id,)).fetchone()
    if not m:
        raise ValueError('Movimento inexistente')
    estado = 'ok' if categoria_id and pessoa_id else 'pendente'
    if m[1] == 'previsto':
        estado = 'previsto'
    ids = [mov_id]
    ids += [r[0] for r in con.execute("SELECT id FROM movimentos WHERE nota_sistema = ?", (f'Taxa da compra #{mov_id}',))]
    if mesma_descricao:
        ids += [r[0] for r in con.execute('''SELECT id FROM movimentos WHERE descricao_norm = ? AND estado = 'pendente'
                                             AND id != ? AND origem = 'extracto' ''', (m[0], mov_id))]
    for i in set(ids):
        nota = 'Classificado pelo utilizador' if i == mov_id or i not in ids[1:] else None
        con.execute('''UPDATE movimentos SET categoria_id = ?, subcategoria_id = ?, pessoa_id = ?,
                       estado = CASE WHEN origem = 'previsto' THEN estado ELSE ? END, aprende = 1,
                       nota_sistema = CASE WHEN nota_sistema LIKE 'Taxa da compra%' THEN nota_sistema
                                           ELSE 'Classificado pelo utilizador' END
                       WHERE id = ?''', (categoria_id, subcategoria_id, pessoa_id, estado, i))
    con.commit()
    return len(set(ids))


def candidatos_previsto(con, mov_id, margem=5):
    m = con.execute('SELECT conta_id, valor, data_lanc FROM movimentos WHERE id = ?', (mov_id,)).fetchone()
    d = dt.date.fromisoformat(m[2])
    return con.execute('''SELECT id, data_lanc, descricao, valor FROM movimentos WHERE origem = 'previsto'
                          AND conta_id = ? AND round(valor, 2) = ? AND data_lanc BETWEEN ? AND ? ORDER BY data_lanc''',
                       (m[0], round(m[1], 2), (d - dt.timedelta(days=margem)).isoformat(),
                        (d + dt.timedelta(days=margem)).isoformat())).fetchall()


def conciliar(con, mov_id, previsto_id):
    """O movimento do banco fica com a classificação do previsto; a descrição escrita vai para as observações."""
    p = con.execute('''SELECT categoria_id, subcategoria_id, pessoa_id, descricao FROM movimentos
                       WHERE id = ? AND origem = 'previsto' ''', (previsto_id,)).fetchone()
    if not p:
        raise ValueError('Movimento previsto inexistente')
    estado = 'ok' if p[0] and p[2] else 'pendente'
    con.execute('''UPDATE movimentos SET categoria_id = ?, subcategoria_id = ?, pessoa_id = ?, estado = ?,
                   observacoes = ?, nota_sistema = 'Conciliado com movimento previsto' WHERE id = ?''',
                (p[0], p[1], p[2], estado, p[3], mov_id))
    con.execute('DELETE FROM movimentos WHERE id = ?', (previsto_id,))
    con.commit()


def criar_previsto(con, conta_id, data, valor, descricao, categoria_id, subcategoria_id, pessoa_id):
    if not descricao or not str(descricao).strip():
        raise ValueError('A descrição é obrigatória')
    if not valor:
        raise ValueError('O valor tem de ser diferente de zero')
    cur = con.execute('''INSERT INTO movimentos(conta_id, data_lanc, descricao, descricao_norm, valor, categoria_id,
                         subcategoria_id, pessoa_id, origem, estado, aprende, criado_em)
                         VALUES (?,?,?,?,?,?,?,?,'previsto','previsto',0,?)''',
                      (conta_id, data, descricao.strip(), norm_desc(descricao), round(valor, 2), categoria_id,
                       subcategoria_id, pessoa_id, _agora()))
    con.commit()
    return cur.lastrowid


def apagar(con, mov_id):
    """Só se apagam movimentos importados, previstos ou manuais; o histórico não se apaga."""
    m = con.execute('SELECT origem FROM movimentos WHERE id = ?', (mov_id,)).fetchone()
    if not m:
        raise ValueError('Movimento inexistente')
    if m[0] == 'historico':
        raise ValueError('Os movimentos do histórico não podem ser apagados; podes corrigir a classificação')
    con.execute('DELETE FROM movimentos WHERE id = ?', (mov_id,))
    con.commit()


def acrescentar(con, tabela, nome, categoria_id=None):
    nome = (nome or '').strip()
    if not nome:
        raise ValueError('O nome não pode ficar vazio')
    if tabela == 'categoria':
        con.execute("INSERT INTO categorias(nome, tipo) VALUES (?, 'custo')", (nome,))
    elif tabela == 'subcategoria':
        con.execute('INSERT INTO subcategorias(categoria_id, nome) VALUES (?,?)', (categoria_id, nome))
    elif tabela == 'pessoa':
        con.execute('INSERT INTO pessoas(nome) VALUES (?)', (nome,))
    else:
        raise ValueError('Lista desconhecida')
    con.commit()


def activar(con, tabela, id_, activa):
    t = {'categoria': 'categorias', 'subcategoria': 'subcategorias', 'pessoa': 'pessoas'}[tabela]
    con.execute(f'UPDATE {t} SET activa = ? WHERE id = ?', (1 if activa else 0, id_))
    if tabela == 'subcategoria' and activa:
        con.execute('UPDATE subcategorias SET na_lista = 1 WHERE id = ?', (id_,))
    con.commit()


def exportar_xlsx(con):
    """R5: todos os movimentos no formato da folha Dados, mais as listas."""
    wb = Workbook(); ws = wb.active; ws.title = 'Dados'
    cab = ['Conta', 'Data Lanc.', 'Data Valor', 'Descrição', 'Valor', 'Categoria', 'Sub-categoria', 'Pessoa',
           'Observações', 'Origem', 'Estado']
    ws.append(cab)
    for c in ws[1]: c.font = Font(bold=True)
    for r in con.execute('''SELECT co.nome, m.data_lanc, m.data_valor, m.descricao, m.valor, c.nome, s.nome, p.nome,
                            m.observacoes, m.origem, m.estado FROM movimentos m JOIN contas co ON co.id = m.conta_id
                            LEFT JOIN categorias c ON c.id = m.categoria_id LEFT JOIN subcategorias s ON s.id = m.subcategoria_id
                            LEFT JOIN pessoas p ON p.id = m.pessoa_id ORDER BY m.data_lanc, m.id'''):
        r = list(r)
        r[1] = dt.date.fromisoformat(r[1]); r[2] = dt.date.fromisoformat(r[2]) if r[2] else None
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = row[2].number_format = 'dd/mm/yyyy'
        row[4].number_format = '#,##0.00 €;-#,##0.00 €'
    for col, w in zip('ABCDEFGHIJK', [18, 12, 12, 48, 12, 34, 26, 10, 30, 10, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    l = wb.create_sheet('Listas')
    l.append(['Categoria', 'Tipo', 'Subcategoria', 'Na lista', 'Activa'])
    for c in l[1]: c.font = Font(bold=True)
    for r in con.execute('''SELECT c.nome, c.tipo, s.nome, s.na_lista, s.activa FROM categorias c
                            LEFT JOIN subcategorias s ON s.categoria_id = c.id ORDER BY c.nome, s.nome'''):
        l.append([r[0], r[1], r[2], 'Sim' if r[3] else ('Não' if r[2] else None), 'Sim' if r[4] else ('Não' if r[2] else None)])
    p = wb.create_sheet('Pessoas')
    p.append(['Pessoa', 'Activa'])
    for r in con.execute('SELECT nome, activa FROM pessoas ORDER BY id'):
        p.append([r[0], 'Sim' if r[1] else 'Não'])
    b = io.BytesIO(); wb.save(b)
    return b.getvalue()
