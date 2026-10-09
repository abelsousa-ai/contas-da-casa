"""Leitura dos extractos bancários para um formato comum.

Cada leitor devolve um Extracto com as linhas pela ordem do ficheiro. Os valores
ficam como vêm do banco; a inversão de sinais do cartão é feita na importação.
"""
import csv, datetime as dt, io, re
from dataclasses import dataclass, field
from openpyxl import load_workbook


class FormatoDesconhecido(Exception):
    pass


@dataclass
class Linha:
    data_lanc: dt.date
    data_valor: dt.date | None
    descricao: str
    valor: float
    saldo: float | None = None
    hora: str | None = None
    tipo: str | None = None      # Revolut: Tipo
    estado: str | None = None    # Revolut: Estado


@dataclass
class Extracto:
    formato: str                 # activobank_conta | activobank_cartao | revolut
    numero: str | None           # número da conta ou do cartão (None na Revolut)
    periodo_de: dt.date | None
    periodo_ate: dt.date | None
    linhas: list = field(default_factory=list)


def _data(v):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    if isinstance(v, str) and v.strip():
        for f in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d'):
            try: return dt.datetime.strptime(v.strip()[:10], f).date()
            except ValueError: pass
    return None


def _num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return round(float(v), 2)
    s = str(v).strip().replace('\u00a0', '').replace(' ', '')
    if ',' in s and '.' in s: s = s.replace('.', '').replace(',', '.')
    elif ',' in s: s = s.replace(',', '.')
    return round(float(s), 2)


def _ler_activobank(conteudo):
    ws = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    if not linhas or not isinstance(linhas[0][0], str):
        raise FormatoDesconhecido('Excel sem cabeçalho reconhecido')
    titulo = linhas[0][0].strip()
    if titulo.upper().startswith('HISTÓRICO DE CONTA NÚMERO'):
        formato = 'activobank_conta'
    elif titulo.upper().startswith('MOVIMENTOS A CRÉDITO NA CONTA CARTÃO NÚMERO'):
        formato = 'activobank_cartao'
    else:
        raise FormatoDesconhecido(f'Excel não reconhecido: "{titulo[:60]}"')
    numero = titulo.split()[-1]
    meta = {str(r[0]).strip().rstrip(':'): r[1] for r in linhas[:12] if r and isinstance(r[0], str)}
    try:
        i_hdr = next(i for i, r in enumerate(linhas) if r and r[0] == 'Data Lanc.')
    except StopIteration:
        raise FormatoDesconhecido('Não encontrei a linha de cabeçalho "Data Lanc."')
    hdr = [str(c).strip() if c else '' for c in linhas[i_hdr]]
    col = {n: hdr.index(n) for n in ('Data Lanc.', 'Data Valor', 'Descrição', 'Valor')}
    col_saldo = hdr.index('Saldo') if 'Saldo' in hdr else None
    ex = Extracto(formato, numero, _data(meta.get('Data de')), _data(meta.get('Data até')))
    for r in linhas[i_hdr + 1:]:
        if not r or r[col['Data Lanc.']] is None: continue
        ex.linhas.append(Linha(
            data_lanc=_data(r[col['Data Lanc.']]), data_valor=_data(r[col['Data Valor']]),
            descricao=str(r[col['Descrição']]).strip(), valor=_num(r[col['Valor']]),
            saldo=_num(r[col_saldo]) if col_saldo is not None else None))
    return ex


REVOLUT_COLUNAS = ['Tipo', 'Produto', 'Data de início', 'Data de Conclusão', 'Descrição',
                   'Montante', 'Comissão', 'Moeda', 'Estado', 'Saldo']


def _ler_revolut(texto):
    leitor = csv.DictReader(io.StringIO(texto))
    if leitor.fieldnames is None or [c.strip() for c in leitor.fieldnames] != REVOLUT_COLUNAS:
        raise FormatoDesconhecido('CSV não reconhecido como extracto da Revolut')
    ex = Extracto('revolut', None, None, None)
    for r in leitor:
        inicio = r['Data de início'].strip()
        conclusao = r['Data de Conclusão'].strip()
        comissao = _num(r['Comissão']) or 0.0
        if r['Moeda'].strip() != 'EUR':
            raise FormatoDesconhecido(f"Moeda {r['Moeda']} não suportada")
        ex.linhas.append(Linha(
            data_lanc=_data(inicio), data_valor=_data(conclusao) if conclusao else None,
            hora=inicio[11:19] or None, descricao=r['Descrição'].strip(),
            valor=_num(r['Montante']), saldo=_num(r['Saldo']),
            tipo=r['Tipo'].strip(), estado=r['Estado'].strip()))
        if comissao:
            # A comissão é um custo à parte do montante (o saldo desce pelos dois)
            ex.linhas.append(Linha(
                data_lanc=_data(inicio), data_valor=_data(conclusao) if conclusao else None,
                hora=inicio[11:19] or None, descricao=f"Comissão Revolut: {r['Descrição'].strip()}",
                valor=-comissao, saldo=None, tipo='Comissão', estado=r['Estado'].strip()))
    datas = [l.data_lanc for l in ex.linhas]
    if datas: ex.periodo_de, ex.periodo_ate = min(datas), max(datas)
    return ex


def ler(nome_ficheiro, conteudo: bytes) -> Extracto:
    nome = nome_ficheiro.lower()
    if nome.endswith(('.xlsx', '.xlsm')):
        return _ler_activobank(conteudo)
    if nome.endswith('.csv'):
        ex = _ler_revolut(conteudo.decode('utf-8-sig'))
        # o nome do ficheiro da Revolut traz o período pedido: account-statement_AAAA-MM-DD_AAAA-MM-DD_...
        m = re.search(r'(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})', nome)
        if m:
            ex.periodo_de, ex.periodo_ate = _data(m.group(1)), _data(m.group(2))
        return ex
    raise FormatoDesconhecido('Só são aceites ficheiros .xlsx (ActivoBank) ou .csv (Revolut)')
